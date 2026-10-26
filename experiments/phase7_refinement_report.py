#!/usr/bin/env python3
"""Publish matched 3-slice ENet refinement figures from three completed evaluations.

Example (run only after all three evaluations complete)::

    python experiments/phase7_refinement_report.py \
      --reports /study/seed-{42,123,2026}/evaluation/report.json \
      --dest /study/report --deck /path/to/source-presentation.pptx

Regenerate figures from the committed completed result without training files::

    python experiments/phase7_refinement_report.py \
      --snapshot notebooks/data/phase7_refinement_results.json \
      --dest results/phase7-report

The result includes a portable patient-level snapshot, CSV tables, PNG/PDF
figures, and optionally a revised copy of the supplied presentation. It never
edits Google Slides. Historical baseline values come from saved notebook
outputs and retain their published precision; no seed values are reconstructed.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import statistics

ROOT = Path(__file__).resolve().parents[1]
ORGANS = ("esophagus", "heart", "trachea", "aorta")
SEEDS = (42, 123, 2026)
PATIENTS = tuple(f"Patient_{n:02d}" for n in (1, 13, 22, 28, 30))
METRICS = ("dice", "hd95_mm", "absolute_rvd")
STAGES = {
    "raw_enet": "3-slice ENet (raw)",
    "raw_refined": "3-slice ENet + refinement (raw)",
    "latest_enet": "3-slice ENet + fixed cleanup",
    "latest_refined": "3-slice ENet + refinement + fixed cleanup",
}
PP = {"min_ml": 1.0, "min_classes": [1, 3], "largest_classes": [4],
      "heart": "unchanged", "closing": False, "fill_holes": False}
INK, BLUE, TEAL, GOLD = "#2e364a", "#7f96ad", "#119d94", "#bb8437"


def read_json(path):
    return json.loads(Path(path).read_text())


def signature(path):
    path = Path(path)
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return {"path": str(path.resolve()), "sha256": digest.hexdigest()}


def save_json(path, value):
    path = Path(path)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, indent=2, allow_nan=False) + "\n")
    temporary.replace(path)


def mean_sd(values):
    values = list(values)
    return {"mean": statistics.mean(values), "sample_sd": statistics.stdev(values)}


def load_reports(paths):
    """Require a complete matched cohort, with traceable frozen checkpoints."""
    reports, rows, sources, selections, histories = {}, [], {}, {}, []
    for path in map(Path, paths):
        report = read_json(path)
        seed = report.get("seed")
        if seed not in SEEDS or seed in reports:
            raise ValueError(f"Expected each of seeds {SEEDS} once: {path}")
        if report.get("status") != "complete":
            raise ValueError(f"Evaluation is not complete: {path}")
        if report.get("epochs") != 25:
            raise ValueError(f"Final refinement publication requires 25 completed epochs: {path}")
        if report.get("metric_source") != "metrics_3d.py" or report.get("postprocessing") != PP:
            raise ValueError(f"Unexpected metric definition or cleanup policy: {path}")
        if set(report.get("patients", {})) != set(PATIENTS):
            raise ValueError(f"Evaluation must contain the five 3-slice ENet validation patients: {path}")
        if not report.get("checkpoint_sha256"):
            raise ValueError(f"Missing ENet checkpoint provenance: {path}")
        protocol_path = path.parent.parent / "protocol.json"
        protocol = read_json(protocol_path)
        if signature(protocol_path)["sha256"] != report.get("protocol_sha256"):
            raise ValueError(f"Protocol hash mismatch: {path}")
        if (protocol.get("seed") != seed or protocol.get("checkpoint_sha256") != report["checkpoint_sha256"]
                or protocol.get("frozen_enet") is not True or protocol.get("refinement_order") != [3, 1, 2]
                or set(protocol.get("validation_patients", [])) != set(PATIENTS)
                or len(protocol.get("train_patients", [])) != 35
                or set(protocol["train_patients"]) & set(PATIENTS)):
            raise ValueError(f"Unmatched 3-slice ENet refinement protocol: {protocol_path}")
        reports[seed] = report
        mode = report.get("checkpoint_selection", protocol.get("checkpoint_selection"))
        if mode != "best" or protocol.get("checkpoint_selection") != mode:
            raise ValueError(f"Checkpoint selection is inconsistent: {path}")
        checkpoints = report.get("refiner_checkpoints", {})
        if set(checkpoints) != {"1", "2", "3"} or any(not c.get("sha256") for c in checkpoints.values()):
            raise ValueError(f"Missing three-organ checkpoint provenance: {path}")
        selections[str(seed)] = {"mode": mode, "metric": protocol.get("checkpoint_selection_metric"),
                                 "organs": {}}
        sources[f"evaluation_seed_{seed}"] = signature(path)
        sources[f"protocol_seed_{seed}"] = signature(protocol_path)
        for patient, stages in report["patients"].items():
            if set(stages) != set(STAGES):
                raise ValueError(f"Missing/unexpected evaluation stages for {patient}: {path}")
            for stage, result in stages.items():
                if set(result.get("organs", {})) != set(ORGANS):
                    raise ValueError(f"Missing organ result: {seed}/{patient}/{stage}")
                for organ, metrics in result["organs"].items():
                    record = {"stage": stage, "seed": seed, "patient": patient, "organ": organ}
                    for metric in METRICS:
                        value = metrics.get(metric)
                        if not isinstance(value, (float, int)) or not math.isfinite(value) or value < 0:
                            raise ValueError(f"Undefined/invalid {metric}: {seed}/{patient}/{organ}/{stage}")
                        if metric == "dice" and value > 1:
                            raise ValueError("Dice cannot exceed 1")
                        record[metric] = float(value)
                    signed_rvd = metrics.get("rvd")
                    if not isinstance(signed_rvd, (int, float)) or not math.isfinite(signed_rvd):
                        raise ValueError("Patient diagnostics require the measured signed RVD")
                    record["rvd"] = float(signed_rvd)
                    if not math.isclose(abs(signed_rvd), record["absolute_rvd"], abs_tol=1e-10):
                        raise ValueError("Absolute RVD must be calculated per patient before averaging")
                    rows.append(record)
        for class_id, organ in ((3, "trachea"), (1, "esophagus"), (2, "heart")):
            model_root = path.parent.parent / "models" / f"class-{class_id}"
            status_path = model_root / "status.json"
            status = read_json(status_path) if status_path.is_file() else {}
            if mode == "best" and (status.get("status") != "complete"
                    or status.get("completed_epochs") != report["epochs"]
                    or status.get("checkpoint_selection") != "best"
                    or not isinstance(status.get("best_epoch"), int)
                    or not 1 <= status["best_epoch"] <= report["epochs"]):
                raise ValueError(f"Best-checkpoint publication requires complete selection metadata: {status_path}")
            selected = {"checkpoint": checkpoints[str(class_id)]}
            if isinstance(status.get("best_epoch"), int):
                selected["selected_epoch"] = status["best_epoch"]
                selected["validation_dice"] = status.get("best_score", status.get("best_validation_dice", status.get("best_dice")))
            selections[str(seed)]["organs"][organ] = selected
            if status_path.is_file():
                sources[f"status_{organ}_seed_{seed}"] = signature(status_path)
            history_path = path.parent.parent / "models" / f"class-{class_id}" / "history.json"
            if history_path.is_file():
                history = read_json(history_path)
                if [row.get("epoch") for row in history] != list(range(1, len(history) + 1)):
                    raise ValueError(f"Training history has missing/repeated epochs: {history_path}")
                if len(history) != report.get("epochs"):
                    raise ValueError(f"Training history does not match completed evaluation budget: {history_path}")
                if mode == "best":
                    validation_scores = [point.get("validation_dice") for point in history]
                    if not all(isinstance(score, (int, float)) and math.isfinite(score)
                               and 0 <= score <= 1 for score in validation_scores):
                        raise ValueError(f"Best-checkpoint publication requires every validation score: {history_path}")
                    best_score = max(validation_scores)
                    best_epoch = validation_scores.index(best_score) + 1
                    if (status["best_epoch"] != best_epoch or not isinstance(status.get("best_score"), (int, float))
                            or not math.isclose(status["best_score"], best_score, rel_tol=0, abs_tol=1e-12)):
                        raise ValueError(f"Selected checkpoint disagrees with the earliest best validation epoch: {status_path}")
                histories.append({"seed": seed, "organ": organ, "points": history})
                sources[f"history_{organ}_seed_{seed}"] = signature(history_path)
            elif mode == "best":
                raise ValueError(f"Best-checkpoint publication requires the complete training history: {history_path}")
    if set(reports) != set(SEEDS):
        raise ValueError(f"All three seeds are required before publishing: {SEEDS}")
    if len({report["epochs"] for report in reports.values()}) != 1:
        raise ValueError("Training budgets differ across seeds")
    if len({value["mode"] for value in selections.values()}) != 1:
        raise ValueError("Checkpoint selection differs across seeds")
    if len({report["checkpoint_sha256"] for report in reports.values()}) != 3:
        raise ValueError("The three seeds must use three distinct ENet checkpoints")
    return rows, sources, selections, histories, reports[42]["epochs"]


def summarize(rows):
    summary = {}
    for stage in STAGES:
        summary[stage] = {}
        for organ in (*ORGANS, "four_organ_mean"):
            summary[stage][organ] = {}
            for metric in METRICS:
                seeds = {str(seed): statistics.mean(row[metric] for row in rows
                    if row["stage"] == stage and row["seed"] == seed
                    and (organ == "four_organ_mean" or row["organ"] == organ)) for seed in SEEDS}
                summary[stage][organ][metric] = {**mean_sd(seeds.values()), "seed_means": seeds}
    deltas = {}
    for first, second in (("raw_enet", "raw_refined"), ("latest_enet", "latest_refined"),
                          ("raw_enet", "latest_enet"), ("raw_refined", "latest_refined")):
        key = f"{second}_minus_{first}"
        deltas[key] = {}
        for organ in (*ORGANS, "four_organ_mean"):
            deltas[key][organ] = {}
            for metric in METRICS:
                seeds = {str(seed): summary[second][organ][metric]["seed_means"][str(seed)]
                         - summary[first][organ][metric]["seed_means"][str(seed)] for seed in SEEDS}
                deltas[key][organ][metric] = {**mean_sd(seeds.values()), "seed_deltas": seeds}
    return summary, deltas


def historical_baseline(notebook, configuration="baseline_0"):
    """Read published summaries; do not turn rounded statistics into seed data."""
    outputs = []
    for cell in read_json(notebook)["cells"]:
        for output in cell.get("outputs", []):
            value = output.get("data", {}).get("text/plain", output.get("text", []))
            outputs.append(value if isinstance(value, str) else "".join(value))
    result = {}
    pair = r"([0-9.]+)\s*\+/-\s*([0-9.]+)"
    for text in outputs:
        for line in text.splitlines():
            prefix = r"\s*\d+\s+" + re.escape(configuration) + r"\s+"
            found = re.match(prefix + r"(esophagus|heart|trachea|aorta)\s+", line)
            if found:
                pairs = re.findall(pair, line)
                if len(pairs) >= 3:
                    result[found[1]] = {metric: {"mean": float(a), "sample_sd": float(b)}
                                       for metric, (a, b) in zip(METRICS, pairs[:3])}
            elif re.match(prefix + r"42,", line):
                pairs = re.findall(pair, line)
                if len(pairs) >= 4:  # 2D Dice then raw metrics; post columns may wrap.
                    result["four_organ_mean"] = {metric: {"mean": float(a), "sample_sd": float(b)}
                                                for metric, (a, b) in zip(METRICS, pairs[1:4])}
    if set(result) != {*ORGANS, "four_organ_mean"}:
        raise ValueError(f"Saved notebook outputs do not contain the complete historical {configuration} summary")
    return {"source": signature(notebook), "configuration": configuration,
            "precision": "As published in saved notebook output; rounded means/SD only",
            "summary": result}


def configure_plotting():
    global plt, np
    import numpy as np
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    font_path = ROOT / "docs/figures/phase7_refinement/fonts/Inter.ttf"
    font_manager.fontManager.addfont(font_path)
    plt.rcParams.update({"font.family": "Inter",
        "font.size": 11, "text.color": INK, "axes.labelcolor": INK, "xtick.color": INK,
        "ytick.color": INK, "axes.edgecolor": "#bbc5ce", "axes.spines.top": False,
        "axes.spines.right": False, "axes.spines.left": False,
        "savefig.facecolor": "white", "pdf.fonttype": 42, "ps.fonttype": 42})


def save_figure(fig, destination, name, manifest):
    fig.savefig(destination / f"{name}.png", dpi=260, bbox_inches="tight")
    fig.savefig(destination / f"{name}.pdf", bbox_inches="tight")
    plt.close(fig)
    manifest[name] = {kind: signature(destination / f"{name}.{kind}") for kind in ("png", "pdf")}


def points(ax, summary, metric, historical=None, annotate=False):
    y = np.arange(len(ORGANS))
    aligned = metric == "dice"
    lower_bounds = []
    methods = [("raw_enet", "3-slice ENet", BLUE), ("raw_refined", "3-slice ENet + refinement", TEAL)]
    if historical:
        methods.insert(0, ("historical", "Original baseline", "#c1c9cf"))
    for index, (stage, label, color) in enumerate(methods):
        values = historical if stage == "historical" else summary[stage]
        means = [values[organ][metric]["mean"] for organ in ORGANS]
        errors = [values[organ][metric]["sample_sd"] for organ in ORGANS]
        lower_bounds.extend(mean - error for mean, error in zip(means, errors))
        positions = y if aligned else y + (index - (len(methods) - 1) / 2) * 0.12
        # A larger open ENet marker keeps both methods visible at the identical
        # aorta score, without moving either point away from its organ row.
        open_marker = aligned and stage == "raw_enet"
        ax.errorbar(means, positions, xerr=errors, fmt="o", color=color, capsize=4,
                    markersize=(9 if open_marker else 6) if aligned else 7,
                    markerfacecolor="white" if open_marker else color,
                    markeredgewidth=1.5 if aligned else 1, label=label, linewidth=1.5)
        if annotate:
            for yy, mean in zip(positions, means):
                ax.annotate(f"{mean:.4f}", (mean, yy), xytext=(0, -16 if stage == "raw_refined" else 9),
                            textcoords="offset points", ha="center", fontsize=9,
                            color="#89959e" if stage == "historical" else color)
    for yy, organ in zip(y, ORGANS):
        if aligned and historical:
            ax.plot([historical[organ][metric]["mean"], summary["raw_enet"][organ][metric]["mean"]],
                    [yy, yy], color="#d7dde2", linewidth=1.2, zorder=0)
        ax.plot([summary[stage][organ][metric]["mean"] for stage in ("raw_enet", "raw_refined")],
                [yy, yy] if aligned else ([yy - 0.06, yy + 0.06] if historical is None else [yy, yy + 0.12]),
                color="#c5d4d1", linewidth=1.2, zorder=0)
    ax.set_yticks(y, [organ.title() for organ in ORGANS])
    ax.set_ylabel("Organ")
    ax.set_ylim(3.55, -.6)
    ax.grid(axis="x", alpha=0.2)
    ax.set_xlabel({"dice": "Patient-level 3D Dice ↑", "hd95_mm": "HD95 (mm) ↓",
                   "absolute_rvd": "Absolute RVD ↓"}[metric])
    if metric == "dice":
        ax.set_xlim(max(0, math.floor((min(lower_bounds) - .05) * 20) / 20), 1.02)
    else:
        ax.set_xlim(left=0)


def figures(snapshot, destination):
    configure_plotting()
    summary, deltas = snapshot["summary"], snapshot["delta_summary"]
    historical = snapshot["historical_baseline"]["summary"]
    manifest = {}
    for name, old in (("slide17_dice", historical), ("refinement_comparison", None)):
        fig, (ax, gain_ax) = plt.subplots(1, 2, figsize=(12.5, 5.5), sharey=True,
                                         gridspec_kw={"width_ratios": [6, 1]})
        points(ax, summary, "dice", historical=old, annotate=True)
        gain_ax.axis("off")
        gain_ax.set_xlim(0, 1)
        reference = old if old is not None else summary["raw_enet"]
        reference_label = "original baseline" if old is not None else "3-slice ENet"
        gain_ax.text(.5, -.43, f"Dice gain\nvs {reference_label}", ha="center", va="center", fontsize=10)
        for yy, organ in enumerate(ORGANS):
            gain = summary["raw_refined"][organ]["dice"]["mean"] - reference[organ]["dice"]["mean"]
            gain_ax.text(.5, yy, f"{gain:+.4f}", ha="center", va="center",
                         fontsize=11, fontweight="bold", color=TEAL)
        fig.legend(*ax.get_legend_handles_labels(), loc="upper center", bbox_to_anchor=(.5, .99),
                   ncol=3, frameon=False, fontsize=10)
        fig.text(.5, .035, "Raw predictions · equal patient means · mean ± sample SD of three seed means", ha="center", fontsize=9)
        fig.text(.5, .005, f"Gain = Dice with refinement − {reference_label} Dice", ha="center", fontsize=9)
        fig.tight_layout(rect=(0, .08, 1, .91))
        save_figure(fig, destination, name, manifest)
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.4))
    for ax, metric, title in zip(axes, ("hd95_mm", "absolute_rvd"), ("Boundary accuracy", "Volume agreement")):
        points(ax, summary, metric, historical)
        ax.set_title(title, fontweight="bold")
    fig.legend(*axes[0].get_legend_handles_labels(), loc="upper center", bbox_to_anchor=(.5, .99),
               ncol=3, frameon=False, fontsize=10)
    fig.text(.5, .01, "Raw predictions · mean ± sample SD of three seed means · original baseline as previously reported", ha="center", fontsize=9)
    fig.tight_layout(rect=(0, .04, 1, .90))
    save_figure(fig, destination, "slide19_other_metrics", manifest)
    fig, axes = plt.subplots(1, 3, figsize=(13, 5.1))
    for ax, metric, title, digits in zip(axes, METRICS, ("Mean Dice ↑", "HD95 (mm) ↓", "Absolute RVD ↓"), (4, 2, 4)):
        first, last = [summary[s]["four_organ_mean"][metric]["mean"] for s in ("raw_enet", "raw_refined")]
        delta = deltas["raw_refined_minus_raw_enet"]["four_organ_mean"][metric]
        ax.axis("off")
        old = historical["four_organ_mean"][metric]["mean"]
        ax.text(.5, .96, title, ha="center", fontsize=17, fontweight="bold")
        for y, label, value, color in ((.79, "Original baseline", old, INK),
                (.52, "3-slice ENet", first, BLUE), (.25, "3-slice ENet + refinement", last, TEAL)):
            ax.text(.5, y, label, ha="center", fontsize=11, color=color)
            ax.text(.5, y-.12, f"{value:.{digits}f}", ha="center", fontsize=21, color=color)
        ax.text(.5, -.015, f"Refinement Δ {delta['mean']:+.{digits}f} ± {delta['sample_sd']:.{digits}f} seed SD",
                ha="center", fontsize=10)
    fig.suptitle("Original baseline → 3-slice ENet → 3-slice ENet + refinement", fontsize=17, fontweight="bold")
    fig.text(.5, .01, "Raw native-volume metrics · same five validation patients and three seeds", ha="center", fontsize=10)
    fig.tight_layout(rect=(0, .04, 1, .93))
    save_figure(fig, destination, "slide18_summary", manifest)
    fig, axes = plt.subplots(1, 3, figsize=(13, 5))
    for ax, metric in zip(axes, METRICS):
        for index, (key, label, color) in enumerate((
                ("latest_enet_minus_raw_enet", "3-slice ENet", BLUE),
                ("latest_refined_minus_raw_refined", "3-slice ENet + refinement", TEAL))):
            values = [deltas[key][o][metric] for o in ORGANS]
            ax.errorbar([v["mean"] for v in values], np.arange(4) + (index-.5)*.15,
                        xerr=[v["sample_sd"] for v in values], fmt="o", capsize=4, color=color, label=label)
        ax.axvline(0, color=INK, linewidth=.7)
        ax.set_yticks(range(4), [o.title() for o in ORGANS]); ax.invert_yaxis()
        ax.set_xlabel("Cleanup Δ " + {"dice": "Dice ↑", "hd95_mm": "HD95 (mm) ↓", "absolute_rvd": "abs RVD ↓"}[metric])
        ax.grid(axis="x", alpha=.2)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(.5, .925),
               ncol=2, frameon=False, fontsize=10)
    fig.suptitle("Effect of the previously chosen cleanup rules", fontsize=17, fontweight="bold")
    fig.text(.5, .045, "Esophagus/trachea: remove components < 1 mL · heart: unchanged · aorta: largest component", ha="center", fontsize=9)
    fig.text(.5, .005, "Fixed-policy evaluation on the same validation patients; no new policy search or held-out estimate", ha="center", fontsize=9)
    fig.tight_layout(rect=(0, .1, 1, .84))
    save_figure(fig, destination, "slide15_fixed_cleanup", manifest)
    fig, ax = plt.subplots(figsize=(12, 5.5)); ax.axis("off")
    steps = [("1  Prepare CT", "HU [−1000, 500]\n1 mm isotropic\n256 × 256 ENet input"),
             ("2  Locate", "Frozen 3-slice ENet\nThree slices\nSame-seed checkpoint"),
             ("3  Correct", "Native CT organ crops\nTrachea → esophagus → heart\nProtect other organ labels"),
             ("4  Measure", "Native-volume predictions\nRaw and fixed cleanup\nDice, HD95 and abs RVD")]
    for i, (title, text) in enumerate(steps):
        x = .02 + i*.25
        ax.text(x, .77, title, transform=ax.transAxes, fontsize=15, fontweight="bold", color=TEAL)
        ax.text(x, .64, text, transform=ax.transAxes, fontsize=11, va="top", linespacing=1.7)
        if i < 3: ax.text(x+.205, .64, "→", transform=ax.transAxes, fontsize=20, color=BLUE)
    ax.text(.5, .17, "Separate organ models and seed runs · aorta keeps the ENet prediction before cleanup", ha="center", fontsize=12)
    ax.text(.5, .06, f"{snapshot['protocol']['epochs']} complete training epochs · 35 training / 5 validation patients · 3 seeds", ha="center", fontsize=11)
    fig.tight_layout()
    save_figure(fig, destination, "slide16_pipeline", manifest)
    plot_histories(snapshot, destination, manifest)
    if __package__:
        from .phase7_refinement_diagnostics import render_diagnostics
    else:
        from phase7_refinement_diagnostics import render_diagnostics
    manifest.update(render_diagnostics(snapshot, destination))
    return manifest


def plot_histories(snapshot, destination, manifest):
    histories = snapshot["training_histories"]
    if not histories:
        return
    fig, axes = plt.subplots(1, 3, figsize=(13, 4.2))
    colors = {42: BLUE, 123: TEAL, 2026: GOLD}
    for ax, organ in zip(axes, ("trachea", "esophagus", "heart")):
        for history in histories:
            if history["organ"] != organ:
                continue
            values = history["points"]
            # The trainer calls the full training objective loss/total depending
            # on format version. Neither quantity is validation Dice.
            loss = [p["loss"].get("total") if isinstance(p.get("loss"), dict)
                    else p.get("loss", p.get("total")) for p in values]
            if all(isinstance(v, (int, float)) and math.isfinite(v) for v in loss):
                ax.plot([p["epoch"] for p in values], loss, label=f"Seed {history['seed']}", color=colors[history["seed"]])
        ax.set(title=organ.title(), xlabel="Completed epoch", ylabel="Training loss")
        ax.grid(alpha=.2)
        if ax.lines: ax.legend(frameon=False, fontsize=9)
    fig.text(.5, .01, "Training objective curves; these are not native-volume validation scores", ha="center", fontsize=9)
    fig.tight_layout(rect=(0, .045, 1, 1))
    save_figure(fig, destination, "refiner_training_curves", manifest)
    # Display the native-volume selection scores recorded during training.
    # Only complete recorded curves are plotted; training loss is never used
    # as a substitute for the checkpoint-selection score.
    validation = []
    for history in histories:
        values = [p.get("validation_dice") for p in history["points"]]
        if all(isinstance(v, (int, float)) and math.isfinite(v) for v in values):
            validation.append((history, values))
    if len(validation) == 9:
        fig, axes = plt.subplots(1, 3, figsize=(13, 4.2))
        for ax, organ in zip(axes, ("trachea", "esophagus", "heart")):
            for history, values in validation:
                if history["organ"] != organ:
                    continue
                seed = history["seed"]
                ax.plot([p["epoch"] for p in history["points"]], values, color=colors[seed], label=f"Seed {seed}")
                selected = snapshot["checkpoint_selection"][str(seed)]["organs"][organ].get("selected_epoch")
                if selected is not None:
                    ax.scatter([selected], [values[selected-1]], color=colors[seed], marker="*", s=90, zorder=3)
            ax.set(title=organ.title(), xlabel="Completed epoch", ylabel="Native-volume validation Dice")
            ax.grid(alpha=.2); ax.legend(frameon=False, fontsize=9)
        fig.text(.5, .01, "Stars: selected organ checkpoints · same five validation patients at every epoch", ha="center", fontsize=9)
        fig.tight_layout(rect=(0, .045, 1, 1))
        save_figure(fig, destination, "refiner_validation_curves", manifest)
    selected_rows = []
    for seed in SEEDS:
        selection = snapshot["checkpoint_selection"][str(seed)]
        for organ in ("trachea", "esophagus", "heart"):
            selected = selection["organs"][organ]
            selected_rows.append([str(seed), organ.title(), selection["mode"], str(selected.get("selected_epoch", "Not recorded"))])
    fig, ax = plt.subplots(figsize=(9, 4.8)); ax.axis("off")
    table = ax.table(cellText=selected_rows, colLabels=["Seed", "Organ", "Selection", "Selected epoch"],
                     cellLoc="center", loc="center")
    table.auto_set_font_size(False); table.set_fontsize(11); table.scale(1, 1.55)
    for (row, _), cell in table.get_celld().items():
        cell.set_edgecolor("#d6dedb")
        if row == 0:
            cell.set_facecolor("#dceae5"); cell.set_text_props(weight="bold", color=INK)
    fig.tight_layout()
    save_figure(fig, destination, "refiner_selected_epochs", manifest)


def update_deck(source, destination, manifest, snapshot):
    """Replace only six affected slides, retaining all other slide XML/content."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    from PIL import Image
    deck = Presentation(source)
    replacements = {15: ("Fixed postprocessing after refinement", "slide15_fixed_cleanup"),
        16: ("3-slice ENet + organ refinement", "slide16_pipeline"),
        17: ("Final model: matched refinement comparison", "slide17_dice"),
        18: ("What refinement changes", "slide18_summary"),
        19: ("Final model: boundary and volume metrics", "slide19_other_metrics"),
        23: ("Organ refinement on the 3-slice ENet checkpoints", "refinement_comparison")}
    expected_titles = {15: "Post-processing", 16: "Pipeline", 17: "Final Model Results",
                       18: "Takeaways", 19: "Final Model", 23: "Organ Refiner"}
    if len(deck.slides) < max(replacements):
        raise ValueError("The supplied presentation does not match the inspected source deck")
    for number, expected in expected_titles.items():
        texts = " ".join(shape.text for shape in deck.slides[number-1].shapes if shape.has_text_frame)
        if expected not in texts:
            raise ValueError(f"Slide {number} changed since inspection; review slide mapping before replacing")
    for number, (title, asset) in replacements.items():
        slide = deck.slides[number-1]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        # Master/background artwork remains. Editable titles and notes accompany
        # high-resolution PNG graphs; vector PDFs are delivered separately.
        title_box = slide.shapes.add_textbox(Pt(48), Pt(23), Pt(635), Pt(44))
        paragraph = title_box.text_frame.paragraphs[0]
        paragraph.text = title
        paragraph.font.name = "Inter Tight"
        paragraph.font.size = Pt(23)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor.from_string(INK[1:])
        image = Path(manifest[asset]["png"]["path"])
        with Image.open(image) as picture:
            ratio = picture.width / picture.height
        width, height = 628, 292
        if width / height > ratio: width = height * ratio
        else: height = width / ratio
        slide.shapes.add_picture(str(image), Pt((720-width)/2), Pt(77+(292-height)/2), width=Pt(width), height=Pt(height))
        footer = slide.shapes.add_textbox(Pt(48), Pt(379), Pt(640), Pt(17))
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = f"Same 5 validation patients · 3 seeds · {snapshot['protocol']['epochs']} refiner training epochs · {number}"
        paragraph.font.name = "Inter"; paragraph.font.size = Pt(8)
        paragraph.font.color.rgb = RGBColor.from_string(INK[1:])
        slide.notes_slide.notes_text_frame.text = (
            f"Updated from completed 3-slice ENet refinement runs. Figure: {asset}.\n"
            "The base ENet checkpoints are frozen, with seed-matched organ refiners. "
            "Esophagus and trachea use batch Dice; heart uses per-image Dice. "
            "Native-volume metrics are averaged equally across five patients within each seed; "
            "error bars are sample SD across three seed means. These are previously used validation patients.\n"
            "Raw graphs compare without cleanup. The fixed-cleanup figure applies the same previously chosen "
            "policy to both pipelines: esophagus/trachea fragments <1 mL, heart unchanged, aorta largest. "
            "This is not a repeated policy search or a new held-out estimate.\n"
            "The original baseline uses published, rounded notebook means/SD; 3-slice ENet and refined values "
            "are recomputed from complete patient-level reports.\n"
            f"Checkpoint selection: {json.dumps(snapshot['checkpoint_selection'], sort_keys=True)}\n"
            "Source and checkpoint provenance: phase7_refinement_results.json and figure_manifest.json."
        )
    output = destination / "phase7-refinement-presentation.pptx"
    deck.save(output)
    # Read-back check also verifies the supplied historical/template slides survive.
    verified = Presentation(output)
    if len(verified.slides) != len(deck.slides):
        raise AssertionError("Presentation slide count changed")
    return {**signature(output), "source": signature(source), "replaced_slides": sorted(replacements),
            "google_slides_modified": False}


def snapshot_from_reports(paths, notebook):
    rows, sources, selections, histories, epochs = load_reports(paths)
    summary, deltas = summarize(rows)
    historical = historical_baseline(notebook)
    historical_phase7 = historical_baseline(notebook, "enet25d_3s_step1_strong_aug_ce_dice_25ep")
    remeasurement = {metric: {
        "historical_reported_mean": historical_phase7["summary"]["four_organ_mean"][metric]["mean"],
        "matched_re_evaluation_mean": summary["raw_enet"]["four_organ_mean"][metric]["mean"],
        "difference": summary["raw_enet"]["four_organ_mean"][metric]["mean"]
                      - historical_phase7["summary"]["four_organ_mean"][metric]["mean"]}
        for metric in METRICS}
    snapshot = {"schema_version": 1, "status": "complete", "protocol": {
        "seeds": list(SEEDS), "validation_patients": list(PATIENTS), "training_patients": 35,
        "epochs": epochs, "frozen_enet": True, "composition_order": ["trachea", "esophagus", "heart"],
        "metric_source": "metrics_3d.py", "space": "native patient geometry",
        "aggregation": "Equal patients within seed, then mean and sample SD across three seed means",
        "validation_status": "Previously inspected validation patients; not an untouched test set",
        "postprocessing": PP, "postprocessing_study": "Fixed prior policy; no new search or held-out estimate"},
        "stages": STAGES, "patient_metrics": rows, "summary": summary, "delta_summary": deltas,
        "checkpoint_selection": selections, "training_histories": histories,
        "sources": sources, "historical_baseline": historical, "historical_phase7": historical_phase7,
        "phase7_re_evaluation": {"metrics": remeasurement,
            "interpretation": "Refinement gains use the re-evaluated same-checkpoint baseline. Historical reported rounded values are retained separately; re-evaluation need not exactly reproduce them."}}
    return snapshot


def load_snapshot(path):
    """Validate the portable final result without requiring archived source paths."""
    snapshot = read_json(path)
    protocol = snapshot.get("protocol", {})
    if (snapshot.get("schema_version") != 1 or snapshot.get("status") != "complete"
            or snapshot.get("stages") != STAGES or protocol.get("seeds") != list(SEEDS)
            or protocol.get("validation_patients") != list(PATIENTS)
            or protocol.get("epochs") != 25 or protocol.get("training_patients") != 35
            or protocol.get("frozen_enet") is not True or protocol.get("postprocessing") != PP
            or protocol.get("metric_source") != "metrics_3d.py"):
        raise ValueError("Expected the completed, matched 25-epoch 3-slice ENet result snapshot")
    rows = snapshot["patient_metrics"]
    keys = [(row["stage"], row["seed"], row["patient"], row["organ"]) for row in rows]
    expected = set(itertools.product(STAGES, SEEDS, PATIENTS, ORGANS))
    if len(keys) != len(expected) or set(keys) != expected:
        raise ValueError("Snapshot must contain each patient/organ/stage/seed exactly once")
    for row in rows:
        for metric in METRICS:
            value = row[metric]
            if not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                raise ValueError(f"Invalid snapshot metric: {metric}")
        if row["dice"] > 1:
            raise ValueError("Dice cannot exceed 1")
        if (not isinstance(row.get("rvd"), (int, float)) or not math.isfinite(row["rvd"])
                or not math.isclose(abs(row["rvd"]), row["absolute_rvd"], rel_tol=0, abs_tol=1e-12)):
            raise ValueError("Snapshot must preserve measured signed RVD and its absolute value")
    summary, deltas = summarize(rows)
    def same_numbers(first, second):
        if isinstance(first, dict):
            return (isinstance(second, dict) and first.keys() == second.keys()
                    and all(same_numbers(value, second[key]) for key, value in first.items()))
        # Python's statistics implementation can differ by a few ulps across versions.
        return isinstance(second, (int, float)) and math.isclose(first, second, rel_tol=0, abs_tol=1e-12)

    if not same_numbers(summary, snapshot["summary"]) or not same_numbers(deltas, snapshot["delta_summary"]):
        raise ValueError("Snapshot summaries disagree with its patient-level measurements")
    histories = snapshot["training_histories"]
    expected_runs = set(itertools.product(SEEDS, ("trachea", "esophagus", "heart")))
    if len(histories) != 9 or {(h["seed"], h["organ"]) for h in histories} != expected_runs:
        raise ValueError("Snapshot must retain all nine refiner histories")
    for history in histories:
        points = history["points"]
        if [point["epoch"] for point in points] != list(range(1, 26)):
            raise ValueError("Snapshot must retain all 25 training epochs")
        scores = [point["validation_dice"] for point in points]
        if any(not isinstance(score, (int, float)) or not math.isfinite(score)
               or not 0 <= score <= 1 for score in scores):
            raise ValueError("Invalid native-volume checkpoint selection scores")
        selection = snapshot["checkpoint_selection"][str(history["seed"])]
        selected = selection["organs"][history["organ"]]
        if (selection["mode"] != "best" or selected["selected_epoch"] != scores.index(max(scores)) + 1
                or not math.isclose(selected["validation_dice"], max(scores), rel_tol=0, abs_tol=1e-12)
                or not selected["checkpoint"].get("sha256")):
            raise ValueError("Snapshot checkpoint selection disagrees with its validation history")
    return snapshot


def main(args):
    snapshot = load_snapshot(args.snapshot) if args.snapshot else snapshot_from_reports(args.reports, args.notebook)
    rows, summary = snapshot["patient_metrics"], snapshot["summary"]
    if args.validate_only:
        print(json.dumps({"status": "valid", "seeds": list(SEEDS), "patient_metric_rows": len(rows),
                          "epochs": snapshot["protocol"]["epochs"]}))
        return
    args.dest.mkdir(parents=True, exist_ok=True)
    save_json(args.dest / "phase7_refinement_results.json", snapshot)
    with (args.dest / "patient_metrics.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["stage", "seed", "patient", "organ", *METRICS, "rvd"])
        writer.writeheader(); writer.writerows(rows)
    with (args.dest / "summary.csv").open("w", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["stage", "organ", "metric", "mean", "sample_sd", "seed_42", "seed_123", "seed_2026"])
        for stage, organ, metric in itertools.product(STAGES, (*ORGANS, "four_organ_mean"), METRICS):
            row = summary[stage][organ][metric]
            writer.writerow([stage, organ, metric, row["mean"], row["sample_sd"], *[row["seed_means"][str(seed)] for seed in SEEDS]])
    manifest = {"source": signature(__file__), "snapshot": signature(args.dest / "phase7_refinement_results.json"),
                "figures": figures(snapshot, args.dest)}
    if args.deck:
        manifest["presentation"] = update_deck(args.deck, args.dest, manifest["figures"], snapshot)
    save_json(args.dest / "figure_manifest.json", manifest)
    print(json.dumps({"status": "complete", "destination": str(args.dest.resolve()), "figures": list(manifest["figures"])}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--reports", type=Path, nargs=3, help="Three completed seed evaluation/report.json files")
    inputs.add_argument("--snapshot", type=Path, help="Portable completed 3-slice ENet result snapshot")
    parser.add_argument("--dest", type=Path, required=True)
    parser.add_argument("--notebook", type=Path, default=ROOT / "notebooks/result_aggrestion.ipynb")
    parser.add_argument("--deck", type=Path, help="Optional downloaded original Google Slides PPTX; writes a revised copy")
    parser.add_argument("--validate-only", action="store_true", help="Validate all results without publishing artifacts")
    main(parser.parse_args())
